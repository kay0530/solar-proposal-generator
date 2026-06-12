"""
ep_estimate.py - EPC estimate / quotation slide (design v2: Institutional Trust Grid)

Professional Japanese estimate form for EPC (purchase) proposals.
Quiet white header (no eyebrow), customer block with navy hairline,
audited-figures table (add_table, 9pt, total row), C_PANEL total band
with orange left bar, notes and issuer signature.

Only included when estimate data is provided (estimate_items in customer_data).
"""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from proposal_generator.utils import (
    CONTENT_BOTTOM, CONTENT_TOP, C_DARK, C_NAVY, C_PANEL, C_ORANGE, C_SUB,
    FONT_BLACK, FONT_BODY, GAP_IN_CARD, MARGIN, SIZE_CAPTION, SIZE_SMALL,
    SIZE_TABLE, SLIDE_W, TABLE_ROW_H,
    add_footer, add_header_bar, add_line, add_multiline_textbox,
    add_number_unit, add_rect, add_section_header, add_table, add_textbox,
    vstack,
)

TITLE = "概算費用お見積書"

TAX_RATE = 0.10
MIN_ITEM_ROWS = 6    # pad with empty rows for a professional look
MAX_ITEM_ROWS = 8    # rows beyond this are merged into a single "その他" line


def _fmt_date(val) -> str:
    """Format date for estimate header."""
    if not val:
        return date.today().strftime("%Y年%m月%d日").replace("年0", "年").replace("月0", "月")
    s = str(val).split(" ")[0]
    m = re.match(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return f"{m.group(1)}年{int(m.group(2))}月{int(m.group(3))}日"
    return s


def _fmt_comma(val) -> str:
    """Format number with commas, no yen symbol."""
    if val is None or val == 0:
        return "-"
    try:
        return f"{int(val):,}"
    except (TypeError, ValueError):
        return str(val)


def _build_estimate_items(data: dict) -> list[dict]:
    """Build estimate line items from customer_data.

    If data contains 'estimate_items' (list of dicts), use that directly.
    Otherwise, auto-build from equipment data.

    Each item: {name, spec, qty, unit, unit_price, amount}
    """
    # Use explicit estimate items if provided
    items = data.get("estimate_items")
    if items:
        return items

    # Auto-build from equipment data
    auto_items = []
    panels = data.get("panels", [])
    for p in panels:
        model = p.get("model", "太陽電池モジュール")
        count = p.get("count", 0) or 0
        watt = p.get("watt_per_unit", 0) or 0
        unit_price = p.get("selling_unit_price", 0) or 0
        if count > 0:
            auto_items.append({
                "name": "太陽電池モジュール",
                "spec": f"{model} ({watt}W)" if watt else model,
                "qty": count,
                "unit": "枚",
                "unit_price": unit_price,
                "amount": int(unit_price * count) if unit_price else None,
            })

    pcs_list = data.get("pcs_list", [])
    for pcs in pcs_list:
        model = pcs.get("model", "パワーコンディショナ")
        count = pcs.get("count", 0) or 0
        kw = pcs.get("kw_per_unit", 0) or 0
        unit_price = pcs.get("selling_unit_price", 0) or 0
        if count > 0:
            auto_items.append({
                "name": "パワーコンディショナ",
                "spec": f"{model} ({kw}kW)" if kw else model,
                "qty": count,
                "unit": "台",
                "unit_price": unit_price,
                "amount": int(unit_price * count) if unit_price else None,
            })

    # Frame & installation
    frame_cost = data.get("estimate_frame_cost", 0) or 0
    if frame_cost > 0:
        auto_items.append({
            "name": "架台・施工費",
            "spec": "",
            "qty": 1,
            "unit": "式",
            "unit_price": frame_cost,
            "amount": frame_cost,
        })

    # Electrical work
    elec_cost = data.get("estimate_electrical_cost", 0) or 0
    if elec_cost > 0:
        auto_items.append({
            "name": "電気工事費",
            "spec": "",
            "qty": 1,
            "unit": "式",
            "unit_price": elec_cost,
            "amount": elec_cost,
        })

    # Battery
    batteries = data.get("batteries", [])
    for bat in batteries:
        model = bat.get("model", "蓄電池")
        count = bat.get("count", 0) or 0
        kwh = bat.get("kwh_per_unit", 0) or 0
        unit_price = bat.get("selling_unit_price", 0) or 0
        if count > 0:
            auto_items.append({
                "name": "蓄電池システム",
                "spec": f"{model} ({kwh}kWh)" if kwh else model,
                "qty": count,
                "unit": "台",
                "unit_price": unit_price,
                "amount": int(unit_price * count) if unit_price else None,
            })

    # Additional items
    extra_items = data.get("estimate_extra_items", [])
    for ex in extra_items:
        auto_items.append(ex)

    return auto_items


def _cap_items(items: list[dict]) -> list[dict]:
    """Merge overflow rows into a single 'その他' line so the audited
    table keeps its exact height math inside the content area."""
    if len(items) <= MAX_ITEM_ROWS:
        return items
    head = items[:MAX_ITEM_ROWS - 1]
    rest = items[MAX_ITEM_ROWS - 1:]
    rest_amount = sum((it.get("amount") or 0) for it in rest)
    head.append({
        "name": "その他",
        "spec": f"{len(rest)}項目一式",
        "qty": 1,
        "unit": "式",
        "unit_price": None,
        "amount": rest_amount if rest_amount > 0 else None,
    })
    return head


def generate(slide, data: dict, logo_path: Path = None) -> None:
    """Render EP_ESTIMATE (EPC quotation page) onto a blank slide."""
    add_header_bar(slide, TITLE, logo_path)

    content_w = SLIDE_W - MARGIN * 2

    company = data.get("company_name", "") or ""
    office = data.get("office_name", "") or ""
    proposal_date = _fmt_date(data.get("proposal_date"))
    tax_display = data.get("tax_display", "税抜")
    opp_id = data.get("opp_id", "")  # reserved for estimate numbering

    # Estimate number from date (fallback: today)
    est_number = data.get("estimate_number", "")
    if not est_number:
        date_part = str(data.get("proposal_date", "")).replace("-", "")[:8]
        if not date_part:
            date_part = date.today().strftime("%Y%m%d")
        est_number = f"EST-{date_part}-001"

    # ---- Estimate items & totals ----
    items = _cap_items(_build_estimate_items(data))

    subtotal = 0
    for item in items:
        amt = item.get("amount")
        if amt and amt > 0:
            subtotal += amt

    # If no item amounts but selling_price exists, use that
    selling_price = data.get("selling_price", 0) or 0
    if subtotal == 0 and selling_price > 0:
        subtotal = selling_price

    tax_amount = int(subtotal * TAX_RATE)
    grand_total = subtotal + tax_amount

    # ---- Table rows (header + items + padding + 3 summary rows) ----
    rows_data = [["No.", "項目", "仕様・詳細", "数量", "単価（円）", "金額（円）"]]
    for i, item in enumerate(items, 1):
        qty_str = f"{item['qty']}{item.get('unit', '')}" if item.get("qty") else ""
        rows_data.append([
            str(i),
            item.get("name", ""),
            item.get("spec", ""),
            qty_str,
            _fmt_comma(item.get("unit_price")),
            _fmt_comma(item.get("amount")),
        ])
    while len(rows_data) < 1 + MIN_ITEM_ROWS:
        rows_data.append(["", "", "", "", "", ""])

    rows_data.append(["", "", "", "", "小計", _fmt_comma(subtotal)])
    rows_data.append(["", "", "", "", "消費税（10%）", _fmt_comma(tax_amount)])
    rows_data.append(["", "", "", "", "合計（税込）", _fmt_comma(grand_total)])
    n_rows = len(rows_data)

    col_widths = [
        Inches(0.45),   # No.
        Inches(2.10),   # Item name
        Inches(3.55),   # Spec/details
        Inches(0.85),   # Qty
        Inches(1.70),   # Unit price
    ]
    col_widths.append(int(content_w) - sum(int(cw) for cw in col_widths))

    # ---- Block heights for vertical justify ----
    cust_h = Inches(0.52)
    subj_h = Inches(0.26)
    table_h = int(TABLE_ROW_H) * n_rows
    band_h = Inches(0.55)
    notes_h = Inches(0.88)
    issuer_h = Inches(0.40)

    ys = vstack(CONTENT_TOP, CONTENT_BOTTOM,
                [cust_h, subj_h, table_h, band_h, notes_h, issuer_h],
                min_gap=GAP_IN_CARD)

    # ---- Customer block: name + navy hairline, date/number right ----
    customer_label = f"{company}"
    if office:
        customer_label += f"  {office}"
    customer_label += "  御中"

    add_textbox(slide, MARGIN, ys[0], Inches(5.8), Inches(0.32),
                customer_label,
                font_name=FONT_BLACK, font_size_pt=14,
                font_color=C_DARK, bold=True)
    add_line(slide, MARGIN, int(ys[0]) + int(Inches(0.42)),
             int(MARGIN) + int(Inches(5.8)), int(ys[0]) + int(Inches(0.42)),
             C_NAVY, width_pt=0.75)

    add_multiline_textbox(
        slide, SLIDE_W - MARGIN - Inches(3.2), ys[0],
        Inches(3.2), Inches(0.44),
        [
            (f"見積日：{proposal_date}", FONT_BODY, SIZE_CAPTION, C_SUB,
             False, PP_ALIGN.RIGHT),
            (f"見積番号：{est_number}", FONT_BODY, SIZE_CAPTION, C_SUB,
             False, PP_ALIGN.RIGHT),
        ],
        line_spacing=1.35)

    # ---- Subject line ----
    capacity = data.get("system_capacity_kw", 0) or 0
    subject = f"太陽光発電設備工事  {capacity:.1f}kW" if capacity else "太陽光発電設備工事"
    add_textbox(slide, MARGIN, ys[1], content_w, subj_h,
                f"件名：{subject}",
                font_name=FONT_BODY, font_size_pt=11,
                font_color=C_DARK, bold=True)

    # ---- Estimate table (audited style, exact height = n_rows * ROW_H) ----
    tbl = add_table(slide, MARGIN, ys[2], content_w, rows_data, col_widths,
                    font_size_pt=SIZE_TABLE, total_row=n_rows - 1)

    # Post-pass: item/spec columns read left; summary labels hug the figure
    for r in range(1, n_rows):
        for c in (1, 2):
            for para in tbl.cell(r, c).text_frame.paragraphs:
                para.alignment = PP_ALIGN.LEFT
        if r >= n_rows - 3:
            for para in tbl.cell(r, 4).text_frame.paragraphs:
                para.alignment = PP_ALIGN.RIGHT

    # ---- Total band: C_PANEL + orange left bar + 20pt number/unit ----
    band_w = Inches(4.8)
    band_x = SLIDE_W - MARGIN - band_w
    band_y = ys[3]
    add_rect(slide, band_x, band_y, band_w, band_h, C_PANEL)
    add_rect(slide, band_x, band_y, Inches(0.05), band_h, C_ORANGE)
    add_textbox(slide, int(band_x) + int(Inches(0.18)),
                int(band_y) + int(Inches(0.17)),
                Inches(1.9), Inches(0.22),
                "お見積り合計（税込）",
                font_name=FONT_BODY, font_size_pt=SIZE_CAPTION,
                font_color=C_SUB, bold=True)
    add_number_unit(slide, int(band_x) + int(Inches(2.0)),
                    int(band_y) + int(Inches(0.08)),
                    int(band_w) - int(Inches(2.2)),
                    int(band_h) - int(Inches(0.16)),
                    _fmt_comma(grand_total) if grand_total > 0 else "—", "円",
                    number_size_pt=20, align=PP_ALIGN.RIGHT)

    # ---- Notes (備考・条件) ----
    validity = data.get("estimate_validity", "本見積書発行日より1ヶ月間")
    delivery = data.get("estimate_delivery", "ご発注後、別途ご相談")
    subsidy_note = data.get("estimate_subsidy_note", "補助金申請費用は別途お見積り")

    notes_y = ys[4]
    add_section_header(slide, MARGIN, notes_y, Inches(4.0), "備考・条件",
                       font_size_pt=11)
    note_lines = [
        (f"見積有効期限：{validity}", FONT_BODY, SIZE_SMALL, C_SUB,
         False, PP_ALIGN.LEFT),
        (f"納期目安：{delivery}", FONT_BODY, SIZE_SMALL, C_SUB,
         False, PP_ALIGN.LEFT),
        (f"備考：{subsidy_note}", FONT_BODY, SIZE_SMALL, C_SUB,
         False, PP_ALIGN.LEFT),
        (f"金額表記：{tax_display}（消費税は税込合計に含む）", FONT_BODY,
         SIZE_SMALL, C_SUB, False, PP_ALIGN.LEFT),
    ]
    add_multiline_textbox(slide, int(MARGIN) + int(Inches(0.20)),
                          int(notes_y) + int(Inches(0.26)),
                          int(content_w) - int(Inches(0.20)),
                          int(notes_h) - int(Inches(0.26)),
                          note_lines, line_spacing=1.35)

    # ---- Issuer signature (right-aligned) ----
    add_multiline_textbox(
        slide, SLIDE_W - MARGIN - Inches(3.5), ys[5],
        Inches(3.5), issuer_h,
        [
            ("株式会社オルテナジー", FONT_BODY, SIZE_CAPTION, C_DARK,
             True, PP_ALIGN.RIGHT),
            ("https://altenergy.co.jp/", FONT_BODY, SIZE_SMALL, C_SUB,
             False, PP_ALIGN.RIGHT),
        ],
        line_spacing=1.35)

    add_footer(slide)
