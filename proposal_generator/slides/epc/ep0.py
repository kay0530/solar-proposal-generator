"""
ep0.py - EPC表紙スライド (Design v2: Institutional Trust Grid)

Full-custom cover on white canvas (no header bar / no orange panel),
mirroring pp0.py:
- 0.18in full-bleed orange band at the very top + color logo top-left
- Left block (cols 0-6): eyebrow -> customer name (+ 御中 inline run)
  -> office/address -> theme title + EPC plan caption -> hairline
  -> date + company
- Right block (cols 8-11): quiet spec list (label + number/unit pairs
  separated by hairlines) — no cards, no fills
- Brand illustration lower-right, bottom hairline + centered copyright
"""
from __future__ import annotations

import re
from pathlib import Path

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from proposal_generator.utils import (
    C_DARK, C_FAINT, C_NAVY, C_ORANGE, C_SUB,
    CONTENT_TOP, FONT_BLACK, FONT_BODY, GAP_BLOCK, MARGIN, SLIDE_W,
    SIZE_CAPTION, SIZE_SMALL,
    add_divider, add_image_contain, add_number_unit, add_rect, add_textbox,
    asset_path, fmt_num, grid_x, grid_w, vstack,
)


def _fmt_date(val) -> str:
    """Format date string: '2026-03-28' or datetime -> '2026年3月28日'."""
    s = str(val).split(" ")[0]
    m = re.match(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return f"{m.group(1)}年{int(m.group(2))}月{int(m.group(3))}日"
    return s


def _yen_parts(v) -> tuple[str, str]:
    """Split a yen amount into (number, unit) for add_number_unit."""
    if v is None:
        return "—", ""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v), ""
    if abs(f) >= 1_0000_0000:
        return f"{f / 1_0000_0000:.2f}", "億円"
    if abs(f) >= 10_000:
        return f"{f / 10_000:,.0f}", "万円"
    return f"{f:,.0f}", "円"


def _add_company_line(slide, x, y, w, h, company: str, name_size_pt: float):
    """Company name + 御中 as two runs in ONE paragraph so long
    (2-line) company names reflow without colliding with 御中."""
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Pt(0)
    tf.margin_right = Pt(0)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.line_spacing = 1.2
    r1 = p.add_run()
    r1.text = company
    r1.font.name = FONT_BLACK
    r1.font.size = Pt(name_size_pt)
    r1.font.color.rgb = C_DARK
    r1.font.bold = True
    r2 = p.add_run()
    r2.text = "　御中"
    r2.font.name = FONT_BODY
    r2.font.size = Pt(16)
    r2.font.color.rgb = C_DARK
    r2.font.bold = False
    return tb


def generate(slide, data: dict, logo_path: Path = None) -> None:
    """
    Render EP0 (EPC cover slide) — design system v2.

    data keys used:
        company_name, office_name, address, proposal_date,
        system_capacity_kw, selling_price,
        subsidy_name, subsidy_amount,
        investment_recovery_yr, annual_gen_kwh
    """
    company = data.get("company_name", "") or ""
    office = data.get("office_name", "") or ""
    address = data.get("address", "") or ""
    prop_date = _fmt_date(data.get("proposal_date", "") or "")
    capacity = data.get("system_capacity_kw")
    selling_price = data.get("selling_price")
    subsidy_name = data.get("subsidy_name", "") or ""
    subsidy_amount = data.get("subsidy_amount")
    recovery_yr = data.get("investment_recovery_yr")
    annual_gen = data.get("annual_gen_kwh")

    # ------------------------------------------------------------------
    # Top brand band + color logo
    # ------------------------------------------------------------------
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.18), C_ORANGE)

    if logo_path and Path(logo_path).exists():
        try:
            add_image_contain(slide, MARGIN, Inches(0.36),
                              Inches(1.60), Inches(0.45), Path(logo_path))
        except Exception:
            pass

    # ------------------------------------------------------------------
    # Left block (cols 0-6) — vertically justified
    # ------------------------------------------------------------------
    lx = grid_x(0)
    lw = grid_w(7)

    # Company name size steps down for long names
    clen = len(company)
    if clen > 25:
        name_size = 22
    elif clen > 18:
        name_size = 26
    else:
        name_size = 30

    blocks: list[tuple] = []

    def draw_eyebrow(y):
        add_textbox(slide, lx, y, lw, Inches(0.22),
                    "御提案書｜自家消費型太陽光発電システム",
                    font_name=FONT_BODY, font_size_pt=SIZE_CAPTION,
                    font_color=C_SUB, bold=True, tracking_pt=1.2)

    blocks.append((Inches(0.24), draw_eyebrow))

    def draw_company(y):
        _add_company_line(slide, lx, y, lw, Inches(1.00), company, name_size)

    blocks.append((Inches(1.00), draw_company))

    if office or address:
        office_h = Inches(0.30) + (Inches(0.22) if address else Inches(0))

        def draw_office(y):
            yy = y
            if office:
                add_textbox(slide, lx, yy, lw, Inches(0.28),
                            office,
                            font_name=FONT_BODY, font_size_pt=14,
                            font_color=C_DARK, bold=True)
                yy += Inches(0.32)
            if address:
                add_textbox(slide, lx, yy, lw, Inches(0.20),
                            f"設置先住所：{address}",
                            font_name=FONT_BODY, font_size_pt=SIZE_CAPTION,
                            font_color=C_SUB)

        blocks.append((office_h, draw_office))

    def draw_theme(y):
        add_textbox(slide, lx, y, lw, Inches(0.45),
                    "自家消費型太陽光発電のご提案",
                    font_name=FONT_BLACK, font_size_pt=24,
                    font_color=C_NAVY, bold=True)
        add_textbox(slide, lx, y + Inches(0.48), lw, Inches(0.20),
                    "EPC（設備購入）方式｜設備購入型の電気代・CO₂削減プラン",
                    font_name=FONT_BODY, font_size_pt=SIZE_CAPTION,
                    font_color=C_SUB, bold=True)

    blocks.append((Inches(0.70), draw_theme))

    def draw_dateline(y):
        add_divider(slide, lx, y, Inches(3.2))
        yy = y + Inches(0.10)
        if prop_date:
            add_textbox(slide, lx, yy, lw, Inches(0.18),
                        prop_date,
                        font_name=FONT_BODY, font_size_pt=SIZE_CAPTION,
                        font_color=C_SUB)
            yy += Inches(0.20)
        add_textbox(slide, lx, yy, lw, Inches(0.18),
                    "株式会社オルテナジー",
                    font_name=FONT_BODY, font_size_pt=SIZE_CAPTION,
                    font_color=C_SUB)

    blocks.append((Inches(0.52) if prop_date else Inches(0.32), draw_dateline))

    left_bottom = Inches(7.30)
    ys = vstack(CONTENT_TOP, left_bottom,
                [h for h, _ in blocks], min_gap=GAP_BLOCK)
    for (h, fn), y in zip(blocks, ys):
        fn(y)

    # ------------------------------------------------------------------
    # Right block (cols 8-11) — quiet spec list, no cards, no fills
    # ------------------------------------------------------------------
    sx = grid_x(8)
    sw = grid_w(4)

    specs: list[tuple] = []
    if capacity:
        specs.append(("設備容量", fmt_num(capacity, 2), "kW"))
    if selling_price:
        num, unit = _yen_parts(selling_price)
        specs.append(("概算投資額（税別）", num, unit))
    if subsidy_name and subsidy_amount:
        num, unit = _yen_parts(subsidy_amount)
        specs.append(("補助金", num, unit))
    if recovery_yr:
        specs.append(("投資回収年", str(recovery_yr), "年"))
    if annual_gen:
        specs.append(("年間想定発電量", fmt_num(annual_gen, 0), "kWh"))

    cover_illust = asset_path("illust_cover.png")
    if specs:
        if cover_illust:
            row_h = Inches(0.80) if len(specs) <= 4 else Inches(0.70)
        else:
            row_h = Inches(0.92)
        total_h = row_h * len(specs)
        if cover_illust:
            # Specs anchored to the top; illustration fills the lower right
            sy = CONTENT_TOP + Inches(0.10)
        else:
            sy = CONTENT_TOP + (left_bottom - CONTENT_TOP - total_h) // 2
        for i, (label, num, unit) in enumerate(specs):
            ry = sy + row_h * i
            add_textbox(slide, sx, ry + Inches(0.06), sw, Inches(0.18),
                        label,
                        font_name=FONT_BODY, font_size_pt=SIZE_CAPTION,
                        font_color=C_SUB, bold=True)
            add_number_unit(slide, sx, ry + Inches(0.26), sw, Inches(0.42),
                            num, unit,
                            number_size_pt=20, unit_size_pt=10,
                            number_color=C_DARK, unit_color=C_SUB,
                            align=PP_ALIGN.LEFT)
            if i < len(specs) - 1:
                add_divider(slide, sx, ry + row_h - Inches(0.08), sw)

    # ------------------------------------------------------------------
    # Cover illustration (lower-right, under the spec list)
    # ------------------------------------------------------------------
    if cover_illust:
        try:
            spec_row_h = Inches(0.80) if len(specs) <= 4 else Inches(0.70)
            il_y = (CONTENT_TOP + Inches(0.10)
                    + spec_row_h * max(len(specs), 1) + Inches(0.25))
            il_h = Inches(7.45) - il_y
            if int(il_h) > int(Inches(1.2)):
                add_image_contain(slide, sx - Inches(0.3), il_y,
                                  sw + Inches(0.3), il_h, cover_illust)
        except Exception:
            pass

    # ------------------------------------------------------------------
    # Bottom: hairline + copyright
    # ------------------------------------------------------------------
    add_divider(slide, MARGIN, Inches(7.70), SLIDE_W - MARGIN * 2)
    add_textbox(slide, MARGIN, Inches(7.80),
                SLIDE_W - MARGIN * 2, Inches(0.20),
                "Copyright 2026 altenergy, Inc.  |  https://altenergy.co.jp/",
                font_name=FONT_BODY, font_size_pt=SIZE_SMALL,
                font_color=C_FAINT, align=PP_ALIGN.CENTER)
