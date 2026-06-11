"""
pp8a.py - ご契約条件サマリー — design system v2

Definition-list layout in two columns: term label (10pt bold navy) +
description (10.5pt, line-spacing 1.35) with hairlines between rows.
Bottom: penalty section — info caption + year table split into two
stacked tables (1-10年 / 11-20年) at 9pt.
"""
from __future__ import annotations

import math
from pathlib import Path

from pptx.util import Inches

from proposal_generator.utils import (
    CONTENT_BOTTOM, CONTENT_TOP, MARGIN, SLIDE_W,
    C_DARK, C_NAVY, C_SUB,
    SIZE_BODY, SIZE_CAPTION, SIZE_SMALL,
    GAP_BLOCK, TABLE_ROW_H,
    add_divider, add_footer, add_header_bar, add_section_header,
    add_table, add_textbox, fmt_yen, vstack,
)

TITLE = "ご契約条件サマリー"
EYEBROW = "04｜ご契約条件"

# Definition-list line metrics
_TERM_H = Inches(0.22)
_DESC_LINE_H = Inches(0.20)   # 10.5pt x 1.35 line spacing
_ITEM_TAIL_H = Inches(0.14)   # divider zone below description
_CHARS_PER_LINE = 32          # full-width chars per line in one column


def _desc_lines(desc: str) -> int:
    return max(1, math.ceil(len(desc) / _CHARS_PER_LINE))


def _item_h(desc: str):
    return _TERM_H + _DESC_LINE_H * _desc_lines(desc) + _ITEM_TAIL_H


def generate(slide, data: dict, logo_path: Path = None) -> None:
    add_header_bar(slide, TITLE, logo_path, eyebrow=EYEBROW)

    years = int(data.get("contract_years", 20) or 20)

    # ---- Definition-list items (【】 stripped, balanced 4 / 4) ----
    left_items = [
        ("設備について",
         "発電設備の施工及びメンテナンスは当社が実施させていただきます。"
         "施工及びメンテナンスの実施のための場所の提供、"
         "その他、契約者にご協力いただくことがあります。"),
        ("契約の期間に関して",
         f"契約期間は{years}年になります。"
         "契約期間満了後は新たに更新契約を締結することにより契約の更新が可能です。"),
        ("契約期間終了後の手続き",
         "契約期間終了後、契約が更新されない場合には太陽光パネル及びPCS等の"
         "主要機器に関しては、弊社の費用負担で撤去いたします。"),
        ("保険に関して",
         "発電設備の火災保険及び賠償責任保険は当社が加入いたします。"
         "当社帰責事由により貴社及び第三者に損害を与えた場合には、"
         "当社が加入する保険により担保される範囲で補償をいたします。"),
    ]
    right_items = [
        ("発電設備の保守",
         "発電設備の定期点検及び故障が発生した場合の補修作業は、"
         "当社の費用負担により実施するものとします。"),
        ("期間中の貴社事由による解約",
         "契約期間中に貴社の都合により解約をされた場合には、"
         "予め設定した違約金が発生いたします。"),
        ("建物等の所有権の移転",
         "建物等の所有権が移転する場合には、新たなる所有者が"
         "オンサイトPPAサービス契約を承継する場合などは"
         "ペナルティが発生しないことといたします。"),
        ("建物等の改修に関して",
         "建物等の改装のために当該設備を一時的に撤去が必要である場合には、"
         "期間及び改修方法を協議の上、貴社の負担で一時的な撤去を行います。"),
    ]

    col_gap = Inches(0.30)
    col_w = (SLIDE_W - MARGIN * 2 - col_gap) / 2
    left_x = MARGIN
    right_x = MARGIN + col_w + col_gap

    def _col_h(items) -> int:
        return sum(int(_item_h(d)) for _, d in items)

    deflist_h = max(_col_h(left_items), _col_h(right_items))
    def_block_h = Inches(0.36) + deflist_h  # section header + list

    # ---- Penalty section data ----
    capacity = float(data.get("system_capacity_kw", 0) or 0)
    selling_price = float(data.get("selling_price", 0) or 0)
    proposal_type = data.get("proposal_type", "ppa")
    depreciation_years = years
    depr_limit = (int(selling_price / depreciation_years)
                  if depreciation_years > 0 else 0)

    milestones = list(range(1, years + 1))
    first_ms = milestones[:10]
    second_ms = milestones[10:20]

    penalty_h = (Inches(0.32)                      # section header
                 + Inches(0.24)                    # info caption
                 + TABLE_ROW_H * 2                 # table 1-10年
                 + (Inches(0.12) + TABLE_ROW_H * 2 if second_ms else Inches(0))
                 + (Inches(0.26) if proposal_type == "ppa" else Inches(0)))

    ys = vstack(CONTENT_TOP, CONTENT_BOTTOM, [def_block_h, penalty_h],
                min_gap=GAP_BLOCK)

    # ---- Definition list ----
    add_section_header(slide, MARGIN, ys[0], SLIDE_W - MARGIN * 2,
                       "オンサイトPPAサービス契約に関して")
    list_y = ys[0] + Inches(0.36)

    def _render_column(x, items):
        iy = list_y
        for i, (term, desc) in enumerate(items):
            add_textbox(slide, x, iy, col_w, _TERM_H,
                        term,
                        font_size_pt=10, font_color=C_NAVY, bold=True)
            desc_h = _DESC_LINE_H * _desc_lines(desc)
            add_textbox(slide, x, iy + _TERM_H, col_w, desc_h,
                        desc,
                        font_size_pt=SIZE_BODY, font_color=C_DARK,
                        line_spacing=1.35)
            iy += int(_item_h(desc))
            if i < len(items) - 1:
                add_divider(slide, x, iy - Inches(0.07), col_w)

    _render_column(left_x, left_items)
    _render_column(right_x, right_items)

    # ---- Penalty section ----
    py = ys[1]
    add_section_header(slide, MARGIN, py, SLIDE_W - MARGIN * 2,
                       "中途解約による違約金に係る設備価額（税抜）")
    py += Inches(0.32)

    equip_label = ("設備価額（PPA事業者負担）" if proposal_type == "ppa"
                   else "設備価額")
    info = (
        f"設備容量 {capacity:.2f}kW　　"
        f"{equip_label} {fmt_yen(selling_price)}　　"
        f"償却年数 {depreciation_years}年　　"
        f"償却限度額 {fmt_yen(depr_limit)}"
    )
    add_textbox(slide, MARGIN, py, SLIDE_W - MARGIN * 2, Inches(0.18),
                info, font_size_pt=SIZE_CAPTION, font_color=C_DARK)
    py += Inches(0.24)

    table_w = SLIDE_W - MARGIN * 2
    label_col_w = Inches(1.30)
    data_col_w = (table_w - label_col_w) / 10  # fixed rhythm for both halves

    def _penalty_rows(ms: list[int]) -> list[list[str]]:
        header = ["経過年数"] + [f"{yr}年" for yr in ms]
        values = ["違約金"]
        for yr in ms:
            if selling_price > 0:
                remaining = max(selling_price - depr_limit * yr, depr_limit)
                values.append(fmt_yen(remaining))
            else:
                values.append("—")
        return [header, values]

    if first_ms:
        w1 = label_col_w + data_col_w * len(first_ms)
        add_table(slide, MARGIN, py, w1, _penalty_rows(first_ms),
                  [label_col_w] + [data_col_w] * len(first_ms),
                  font_size_pt=SIZE_CAPTION)
        py += TABLE_ROW_H * 2
    if second_ms:
        py += Inches(0.12)
        w2 = label_col_w + data_col_w * len(second_ms)
        add_table(slide, MARGIN, py, w2, _penalty_rows(second_ms),
                  [label_col_w] + [data_col_w] * len(second_ms),
                  font_size_pt=SIZE_CAPTION)
        py += TABLE_ROW_H * 2

    if proposal_type == "ppa":
        add_textbox(slide, MARGIN, py + Inches(0.06),
                    SLIDE_W - MARGIN * 2, Inches(0.18),
                    "※ 上記は設備の残存簿価に基づく概算違約金です",
                    font_size_pt=SIZE_SMALL, font_color=C_SUB)

    add_footer(slide)
