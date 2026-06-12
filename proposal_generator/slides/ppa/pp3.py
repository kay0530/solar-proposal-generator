"""
pp3.py - 導入メリット（static スライド）

Design v2 "Institutional Trust Grid":
  4 merit cards in a 2x2 grid (grid cols 0-5 / 6-11 per row).
  Each card: eyebrow tag + 14pt headline (numbers as orange runs)
  + hairline divider + 10.5pt body. No pseudo-KPI glyphs.
"""
from __future__ import annotations

from pathlib import Path

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from proposal_generator.utils import (
    add_icon,
    CONTENT_BOTTOM, CONTENT_TOP, C_DARK, C_ORANGE,
    FONT_BLACK, FONT_BODY,
    LINE_SPACING_BODY, SIZE_BODY, SIZE_CAPTION, SIZE_H2,
    add_card_with_accent, add_divider, add_footer, add_header_bar,
    add_textbox, grid_x, grid_w, vstack,
)

TITLE = "導入メリット"
EYEBROW = "01｜導入の背景"

# headline: list of (text, is_number) — number segments render orange
MERITS = [
    {
        "tag": "初期費用", "icon": "zero_yen",
        "headline": [("初期費用 ", False), ("0円", True)],
        "body": "太陽光パネル・PCSなど全設備の設置費用はPPA事業者が負担。"
                "お客様の初期投資なしで太陽光発電を導入できます。",
    },
    {
        "tag": "料金", "icon": "yen",
        "headline": [("長期固定単価で電気代削減", False)],
        "body": "契約期間中のPPA単価は固定。現行の電気料金より安い単価で"
                "利用でき、電力市場の価格高騰リスクを回避できます。",
    },
    {
        "tag": "保守", "icon": "check",
        "headline": [("維持管理の手間なし", False)],
        "body": "設備の保守・点検・保険対応はすべてPPA事業者が実施。"
                "24時間モニタリングで安定稼働を支えます。",
    },
    {
        "tag": "環境", "icon": "leaf",
        "headline": [("CO₂削減・脱炭素経営", False)],
        "body": "再生可能エネルギーによる発電でCO₂排出量を大幅に削減。"
                "蓄電池との組み合わせで停電時のBCP対応も強化できます。",
    },
]


def _add_headline(slide, x, y, w, h, segments):
    """14pt bold headline in one paragraph; number segments as orange runs."""
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(0)
    tf.margin_right = Pt(0)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    for text, is_number in segments:
        run = p.add_run()
        run.text = text
        run.font.name = FONT_BLACK
        run.font.size = Pt(SIZE_H2)
        run.font.bold = True
        run.font.color.rgb = C_ORANGE if is_number else C_DARK
    return tb


def generate(slide, data: dict, logo_path: Path = None) -> None:
    add_header_bar(slide, TITLE, logo_path, eyebrow=EYEBROW)

    card_h = Inches(2.45)
    ys = vstack(CONTENT_TOP, CONTENT_BOTTOM, [card_h, card_h])

    for i, merit in enumerate(MERITS):
        col = i % 2
        row = i // 2
        x = grid_x(col * 6)
        w = grid_w(6)
        cx, cy, cw, ch = add_card_with_accent(slide, x, ys[row], w, card_h)

        add_textbox(slide, cx, cy + Inches(0.08), cw, Inches(0.20),
                    merit["tag"],
                    font_name=FONT_BODY, font_size_pt=SIZE_CAPTION,
                    font_color=C_ORANGE, bold=True, tracking_pt=1.2)
        if merit.get("icon"):
            add_icon(slide, merit["icon"],
                     cx + cw - Inches(0.50), cy + Inches(0.06), Inches(0.44))

        _add_headline(slide, cx, cy + Inches(0.40), cw, Inches(0.34),
                      merit["headline"])

        add_divider(slide, cx, cy + Inches(0.84), cw)

        add_textbox(slide, cx, cy + Inches(0.98), cw, ch - Inches(1.04),
                    merit["body"],
                    font_name=FONT_BODY, font_size_pt=SIZE_BODY,
                    font_color=C_DARK, line_spacing=LINE_SPACING_BODY)

    add_footer(slide)
