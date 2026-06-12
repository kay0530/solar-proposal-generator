"""
new_competitor.py - 他社比較表スライド (design v2: Institutional Trust Grid)

Audited comparison table (オルテナジー vs 他社A/他社B):
  - v2 add_table with highlight_col on the オルテナジー column
    (C_TINT body + orange header emphasis), white header + navy rule, 9pt
  - Three strength cards (flat line icons) restating the differentiators
  - Source note bottom-right
Rows come from data["competitors"] (list of lists) or the
proposal-type-specific defaults; all rows are always rendered.
"""
from __future__ import annotations

from pathlib import Path

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from proposal_generator.utils import (
    CONTENT_BOTTOM, CONTENT_TOP, C_DARK, C_ORANGE, C_SUB,
    GAP_CARD, MARGIN, SIZE_BODY, SIZE_CAPTION, SIZE_SMALL, SIZE_TABLE,
    SLIDE_W, TABLE_ROW_H,
    add_card_with_accent, add_footer, add_header_bar, add_icon,
    add_section_header, add_table, add_textbox, grid_w, grid_x, vstack,
)

TITLE = "他社比較"
EYEBROW = "02｜比較検討"

_OURS = "オルテナジー（当社）"

# Default comparison data (used when data.get("competitors") is not available)
_DEFAULT_ROWS_PPA = [
    ["比較項目",     _OURS,                              "他社A",                  "他社B"],
    ["PPA単価",      "業界最安水準",                     "標準的",                 "やや高め"],
    ["サービス内容", "設計〜施工〜運用まで一貫対応",     "設計・施工のみ",         "施工・運用"],
    ["メンテナンス", "24時間遠隔監視＋定期点検込み",     "オプション（別途費用）", "年1回点検のみ"],
    ["契約柔軟性",   "契約期間・条件をカスタマイズ可能", "固定プランのみ",         "一部カスタマイズ可"],
    ["実績",         "全国500件以上の導入実績",          "関東中心100件程度",      "50件程度"],
]
_DEFAULT_ROWS_EPC = [
    ["比較項目",     _OURS,                              "他社A",                  "他社B"],
    ["kW単価",       "業界最安水準",                     "標準的",                 "やや高め"],
    ["サービス内容", "設計〜施工〜運用まで一貫対応",     "設計・施工のみ",         "施工・運用"],
    ["メンテナンス", "24時間遠隔監視＋定期点検込み",     "オプション（別途費用）", "年1回点検のみ"],
    ["保証内容",     "パネル25年・PCS10年・施工10年",    "パネル25年・PCS5年",     "パネル25年のみ"],
    ["実績",         "全国500件以上の導入実績",          "関東中心100件程度",      "50件程度"],
]

# Strength cards: (icon, title, description) — restate the オルテナジー column
_STRENGTHS = [
    ("panel", "設計〜運用まで一貫対応",  "ワンストップ体制でトータルコストを最適化"),
    ("check", "メンテナンス込み",        "24時間遠隔監視と定期点検を標準で提供"),
    ("doc",   "全国500件以上の導入実績", "業種・規模を問わない豊富な導入ノウハウ"),
]


def _resolve_rows(data: dict) -> list[list]:
    """Return comparison rows: custom data if well-formed, else defaults."""
    competitors = data.get("competitors")
    if (isinstance(competitors, list) and competitors
            and all(isinstance(r, (list, tuple)) and len(r) >= 2
                    for r in competitors)
            and len({len(r) for r in competitors}) == 1):
        return [list(r) for r in competitors]
    return (_DEFAULT_ROWS_EPC if data.get("proposal_type") == "epc"
            else _DEFAULT_ROWS_PPA)


def generate(slide, data: dict, logo_path: Path = None) -> None:
    add_header_bar(slide, TITLE, logo_path, eyebrow=EYEBROW)
    content_w = SLIDE_W - MARGIN * 2

    rows = _resolve_rows(data)
    n_rows = len(rows)
    n_cols = len(rows[0])

    # ---- Block heights for vertical justify ----
    lead_h = Inches(0.26)
    table_block_h = int(Inches(0.36)) + int(TABLE_ROW_H) * n_rows
    cards_block_h = int(Inches(0.36)) + int(Inches(1.00))
    note_h = Inches(0.20)

    ys = vstack(CONTENT_TOP, CONTENT_BOTTOM,
                [lead_h, table_block_h, cards_block_h, note_h],
                min_gap=GAP_CARD)

    # ---- Lead ----
    add_textbox(slide, MARGIN, ys[0], content_w, lead_h,
                "PPA単価・サービス・実績など、主要項目で他社と比較しました。",
                font_size_pt=SIZE_BODY, font_color=C_SUB)

    # ---- Audited comparison table ----
    add_section_header(slide, MARGIN, ys[1], content_w, "主要項目での他社比較")
    label_w = int(Inches(1.7))
    rest_w = (int(content_w) - label_w) // (n_cols - 1)
    col_widths = [label_w] + [rest_w] * (n_cols - 1)

    table_y = int(ys[1]) + int(Inches(0.36))
    tbl = add_table(slide, MARGIN, table_y, content_w, rows, col_widths,
                    font_size_pt=SIZE_TABLE, highlight_col=1)
    # Orange emphasis on the オルテナジー header cell
    try:
        for para in tbl.cell(0, 1).text_frame.paragraphs:
            for run in para.runs:
                run.font.color.rgb = C_ORANGE
    except Exception:
        pass

    # ---- Strength cards ----
    cards_y = int(ys[2])
    add_section_header(slide, MARGIN, cards_y, content_w,
                       "オルテナジーが選ばれる理由")
    card_y = cards_y + int(Inches(0.36))
    card_h = Inches(1.00)
    icon_s = Inches(0.40)
    for i, (icon, title, desc) in enumerate(_STRENGTHS):
        cx, cy, cw, ch = add_card_with_accent(
            slide, grid_x(i * 4), card_y, grid_w(4), card_h,
            accent_position="left")
        add_icon(slide, icon, cx,
                 int(cy) + (int(ch) - int(icon_s)) // 2, size=icon_s)
        tx = int(cx) + int(icon_s) + int(Inches(0.14))
        tw = int(cw) - int(icon_s) - int(Inches(0.14))
        add_textbox(slide, tx, int(cy) + int(Inches(0.06)), tw, Inches(0.22),
                    title, font_size_pt=SIZE_BODY, font_color=C_DARK,
                    bold=True)
        add_textbox(slide, tx, int(cy) + int(Inches(0.32)), tw, Inches(0.40),
                    desc, font_size_pt=SIZE_CAPTION, font_color=C_SUB,
                    line_spacing=1.35)

    # ---- Source note ----
    add_textbox(slide, MARGIN, ys[3], content_w, note_h,
                "※ 他社情報は一般的な市場調査に基づく参考値です",
                font_size_pt=SIZE_SMALL, font_color=C_SUB,
                align=PP_ALIGN.RIGHT)

    add_footer(slide)
